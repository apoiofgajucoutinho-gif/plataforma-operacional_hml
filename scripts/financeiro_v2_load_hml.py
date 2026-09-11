
import argparse
import json
import os
import sys
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

TENANT_ID = "ff24d2bc-22e2-4a5b-bc8c-f6d25a7fa3b0"
HML_REF = "oerdsmgiebquecqwcbox"
SOURCE_FILE = Path(r"D:\Jass_\Download_HD\Fluxo de caixa V2.xlsx")
SOURCE_SHEET = "Fluxo de caixa V2"
SNAPSHOT_KEY = "financeiro_v2_preload_20260911"
BATCH_SIZE = 200

CATEGORY_DEFS = [
    ("entrada", "Receitas", "receitas", "operacional", "nao_aplicavel", ["Hotmart", "Greenn", "TMB", "Clínica", "Palestras", "Venda por fora", "Outras receitas"]),
    ("saida", "Pessoas", "pessoas", "operacional", "fixo", ["Salários", "Pró-labore", "Encargos", "Comissões", "Prestadores recorrentes"]),
    ("saida", "Marketing & Aquisição", "marketing_aquisicao", "operacional", "variavel", ["Meta Ads", "Mídia paga", "Agência / tráfego", "Aquisição"]),
    ("saida", "Conteúdo & Crescimento", "conteudo_crescimento", "operacional", "variavel", ["Produção de conteúdo", "Edição", "Teleprompter", "Sessão de fotos", "Estrutura", "Branding", "Outros"]),
    ("saida", "Tecnologia & Ferramentas", "tecnologia_ferramentas", "operacional", "fixo", ["Canva", "ChatGPT", "ActiveCampaign", "HostGator", "Cademi", "Microsoft", "Notion", "Ferramentas/SaaS", "Automação / CRM", "Outras ferramentas"]),
    ("saida", "Pedagógico / Entrega", "pedagogico_entrega", "operacional", "variavel", ["Professoras convidadas", "Certificado MEC", "Materiais", "Entrega dos cursos"]),
    ("saida", "Operação Clínica", "operacao_clinica", "operacional", "variavel", ["Repasses", "Materiais clínicos", "Alimentação/encontro", "Despesas clínicas", "Transporte / deslocamento"]),
    ("saida", "Administrativo", "administrativo", "operacional", "fixo", ["Contabilidade", "Internet", "Telefone", "Conselho profissional", "Despesas administrativas"]),
    ("saida", "Tributos & Taxas", "tributos_taxas", "operacional", "variavel", ["DAS/DASS", "DARF", "IOF", "Taxas bancárias", "Imposto", "Outros tributos"]),
    ("saida", "Financeiro / Patrimonial", "financeiro_patrimonial", "nao_operacional", "nao_aplicavel", ["Aplicação financeira", "Resgate", "Movimentação financeira"]),
    ("saida", "Distribuição aos Sócios", "distribuicao_socios", "nao_operacional", "nao_aplicavel", ["Distribuição de lucros"]),
    ("saida", "Outros", "outros", "operacional", "nao_aplicavel", ["Outros"]),
]
CENTERS = ["Infoprodutos", "Clínica", "Palestras", "Administrativo / Corporativo", "Não operacional"]


def read_env():
    env = {}
    for file_name in [".env.local", ".env"]:
        p = Path(file_name)
        if not p.exists():
            continue
        for line in p.read_text(encoding="utf-8", errors="ignore").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            env.setdefault(k.strip(), v.strip().strip('"').strip("'"))
    env.update(os.environ)
    return env


def require_env():
    env = read_env()
    url = env.get("NEXT_PUBLIC_SUPABASE_URL") or env.get("SUPABASE_URL")
    key = env.get("SUPABASE_SERVICE_ROLE_KEY") or env.get("SUPABASE_SERVICE_KEY")
    if not url or not key:
        raise SystemExit("Missing Supabase URL/service key in environment.")
    if HML_REF not in url:
        raise SystemExit(f"Refusing to run: Supabase URL is not HML ref {HML_REF}.")
    return url.rstrip("/"), key


def request(method, path, body=None, params=None):
    url, key = require_env()
    endpoint = f"{url}/rest/v1/{path}"
    if params:
        endpoint += "?" + urllib.parse.urlencode(params, doseq=True)
    data = None if body is None else json.dumps(body, default=str).encode("utf-8")
    headers = {
        "apikey": key,
        "Authorization": f"Bearer {key}",
        "Content-Type": "application/json",
        "Prefer": "return=representation,resolution=merge-duplicates",
    }
    req = urllib.request.Request(endpoint, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            text = resp.read().decode("utf-8")
            return json.loads(text) if text else []
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="ignore")
        raise RuntimeError(f"{method} {path} failed {exc.code}: {detail}") from exc


def get_all(table, select="*", extra=None):
    out = []
    start = 0
    while True:
        params = {"select": select, "tenant_id": f"eq.{TENANT_ID}", "offset": str(start), "limit": "1000"}
        if extra:
            params.update(extra)
        batch = request("GET", table, params=params)
        out.extend(batch)
        if len(batch) < 1000:
            return out
        start += 1000


def money(v):
    if v is None or v == "":
        return None
    d = Decimal(str(v)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return d


def clean(value):
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def norm(value):
    import unicodedata
    text = clean(value).lower()
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    return text


def iso_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raise ValueError(f"Invalid date: {value!r}")


def month_start(iso):
    return iso[:7] + "-01"


def read_sheet():
    wb = load_workbook(SOURCE_FILE, data_only=True)
    ws = wb[SOURCE_SHEET]
    headers = [clean(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))[0:9]]
    expected = ["Data", "Tipo", "Natureza", "Categoria", "Subcategoria", "Centro de Resultado", "Descricao", "Pagamento", "Valor"]
    canonical_headers = [h.replace("ção", "cao").replace("Descrição", "Descricao") for h in headers]
    rows = []
    ignored = []
    for ix, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        cells = list(row[:9])
        if not any(cells):
            continue
        amount = money(cells[8])
        if amount is None or amount == 0:
            ignored.append({"row": ix, "reason": "sem_valor_monetario"})
            continue
        row_data = dict(zip(expected, cells))
        row_data["source_row"] = ix
        row_data["Valor"] = amount
        row_data["Data"] = iso_date(row_data["Data"])
        rows.append(row_data)
    return rows, ignored, headers, canonical_headers


def classify(row):
    tipo = "entrada" if norm(row["Tipo"]) == "entrada" else "saida"
    n_cat, n_sub, n_desc, n_center, n_nat = map(norm, [row["Categoria"], row["Subcategoria"], row["Descricao"], row["Centro de Resultado"], row["Natureza"]])
    status = "trusted"
    notes = []
    if tipo == "entrada":
        category, subcategory = "Receitas", "Outras receitas"
        if "hotmart" in n_desc or "hotmart" in n_sub: subcategory = "Hotmart"
        elif "greenn" in n_desc or "greenn" in n_sub: subcategory = "Greenn"
        elif "tmb" in n_desc or "tbm" in n_desc or "tmb" in n_sub or "tbm" in n_sub: subcategory = "TMB"
        elif "clinica" in n_cat or "clinica" in n_sub: subcategory = "Clínica"
        elif "palestra" in n_cat or "palestra" in n_sub: subcategory = "Palestras"
        elif "venda por fora" in n_sub: subcategory = "Venda por fora"
    elif "distribuicao" in n_sub or "lucros" in n_sub:
        category, subcategory = "Distribuição aos Sócios", "Distribuição de lucros"
    elif n_cat == "pessoal":
        category = "Pessoas"
        subcategory = "Pró-labore" if "pro" in n_sub else "Encargos" if "inss" in n_sub else "Salários" if "funcionario" in n_sub else "Prestadores recorrentes"
    elif n_cat == "marketing":
        category, subcategory = "Marketing & Aquisição", "Mídia paga"
    elif n_cat == "crescimento":
        category = "Conteúdo & Crescimento"
        if "branding" in n_sub: subcategory = "Branding"
        elif "teleprompter" in n_desc or "teleprompter" in n_sub: subcategory = "Teleprompter"
        elif "estrutura" in n_sub: subcategory = "Estrutura"
        else: subcategory = "Outros"
    elif n_cat == "infoproduto" and ("activecamp" in n_desc or "cademi" in n_desc):
        category = "Tecnologia & Ferramentas"
        subcategory = "ActiveCampaign" if "activecamp" in n_desc else "Cademi"
    elif n_cat == "infoproduto" and "funcionario" in n_sub:
        category, subcategory = "Pessoas", "Prestadores recorrentes"
    elif n_cat == "administrativo" and ("ferramenta" in n_sub or "canva" in n_desc or "cloude" in n_desc):
        category = "Tecnologia & Ferramentas"
        if "canva" in n_desc: subcategory = "Canva"
        elif "cloude" in n_desc:
            subcategory = "Outras ferramentas"; status = "review"; notes.append("CLOUDE mantido para validacao semantica; nao normalizado para Claude")
        else: subcategory = "Ferramentas/SaaS"
    elif n_cat == "administrativo":
        category = "Administrativo"
        if "contabilidade" in n_sub: subcategory = "Contabilidade"
        elif "internet" in n_desc: subcategory = "Internet"
        elif "chip" in n_desc or "vivo" in n_desc or "telefone" in n_desc: subcategory = "Telefone"
        elif "conselho" in n_sub or "academia" in n_sub: subcategory = "Conselho profissional"
        else: subcategory = "Despesas administrativas"
    elif n_cat == "impostos":
        category, subcategory = "Tributos & Taxas", "Imposto"
    elif n_cat == "financeiro":
        if "taxa" in n_sub:
            category, subcategory = "Tributos & Taxas", "Taxas bancárias"
        elif "troca de dinheiro" in n_desc:
            category, subcategory = "Financeiro / Patrimonial", "Movimentação financeira"
        else:
            category, subcategory = "Financeiro / Patrimonial", "Aplicação financeira"
    elif n_cat == "pedagogico":
        category = "Pedagógico / Entrega"
        subcategory = "Certificado MEC" if "certificado" in n_sub else "Professoras convidadas" if "professor" in n_sub else "Materiais" if not n_sub and ("360imprimir" in n_desc or "elo7" in n_desc) else "Entrega dos cursos"
    elif n_cat == "clinica":
        category = "Operação Clínica"
        if "transporte" in n_sub or "deslocamento" in n_sub: subcategory = "Transporte / deslocamento"
        elif "comida" in n_sub or "aliment" in n_sub: subcategory = "Alimentação/encontro"
        else: subcategory = "Despesas clínicas"
    else:
        category, subcategory = "Outros", "Outros"; status = "review"; notes.append("categoria_original_sem_mapeamento_canonico")

    if not clean(row["Subcategoria"]) and not ("360imprimir" in n_desc or "elo7" in n_desc or "troca de dinheiro" in n_desc):
        status = "review"; notes.append("subcategoria_ausente_na_planilha")

    center = "Infoprodutos"
    if "clinica" in n_center: center = "Clínica"
    elif "palestra" in n_center: center = "Palestras"
    elif "nao operacional" in n_center: center = "Não operacional"
    elif "administrativo" in n_center: center = "Administrativo / Corporativo"

    if norm(row["Natureza"]) == "fixo":
        fluxo, comportamento = "operacional", "fixo"
    elif norm(row["Natureza"]) == "variavel":
        fluxo, comportamento = "operacional", "variavel"
    elif "nao operacional" in norm(row["Natureza"]):
        fluxo, comportamento = "nao_operacional", "nao_aplicavel"
    else:
        fluxo, comportamento = "operacional", "nao_aplicavel"
    if category in ("Financeiro / Patrimonial", "Distribuição aos Sócios") or center == "Não operacional":
        fluxo, comportamento = "nao_operacional", "nao_aplicavel"
    return {"tipo": tipo, "category": category, "subcategory": subcategory, "center": center, "natureza_fluxo": fluxo, "comportamento": comportamento, "classificacao_status": status, "notes": notes}


def upsert(table, rows, conflict):
    if not rows:
        return []
    params = {"on_conflict": conflict}
    return request("POST", table, rows, params=params)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    rows, ignored, headers, canonical_headers = read_sheet()
    classifications = [classify(r) for r in rows]
    totals = defaultdict(Decimal)
    for r, c in zip(rows, classifications):
        totals[c["tipo"]] += r["Valor"]
    report = {
        "source_file": str(SOURCE_FILE), "headers": headers, "rows_total": len(rows) + len(ignored), "monetary_valid_rows": len(rows), "ignored_rows": ignored,
        "totals": {k: str(v) for k, v in totals.items()}, "classification_status": dict(Counter(c["classificacao_status"] for c in classifications)),
        "categories": dict(Counter(c["category"] for c in classifications)), "subcategories": dict(Counter(c["subcategory"] for c in classifications)), "centers": dict(Counter(c["center"] for c in classifications)),
    }
    if not args.apply:
        print(json.dumps({"mode":"dry-run", **report}, ensure_ascii=False, indent=2))
        return

    existing = get_all("fin_lancamentos", select="*")
    if existing:
        snapshots = [{"tenant_id": TENANT_ID, "snapshot_key": SNAPSHOT_KEY, "lancamento_id": e["id"], "row_data": e, "created_by": "financeiro_v2_load_hml"} for e in existing]
        for i in range(0, len(snapshots), BATCH_SIZE):
            upsert("fin_lancamentos_snapshot", snapshots[i:i+BATCH_SIZE], "tenant_id,snapshot_key,lancamento_id")
    request("DELETE", "fin_lancamentos", params={"tenant_id": f"eq.{TENANT_ID}"})

    # Upsert minimal canonical cadastros. Existing legacy cadastros are preserved but no longer used by loaded rows.
    centers = [{"tenant_id": TENANT_ID, "nome": name, "ativo": True} for name in CENTERS]
    upsert("fin_centros_resultado", centers, "tenant_id,nome")
    bank = upsert("fin_bancos", [{"tenant_id": TENANT_ID, "nome": "Conta bancária principal", "apelido": "Caixa", "saldo_inicial": 0, "ativo": True}], "tenant_id,nome")[0]
    card = upsert("fin_cartoes", [{"tenant_id": TENANT_ID, "banco_id": bank["id"], "nome": "Cartão corporativo", "dia_fechamento": 5, "dia_vencimento": 12, "limite": None, "ativo": True}], "tenant_id,nome")[0]
    cats = []
    for tipo, name, dre, fluxo, comp, subs in CATEGORY_DEFS:
        cats.append({"tenant_id": TENANT_ID, "tipo": tipo, "nome": name, "dre_grupo": dre, "natureza_fluxo_padrao": fluxo, "comportamento_padrao": comp, "ativo": True})
    cat_rows = upsert("fin_categorias", cats, "tenant_id,tipo,nome")
    cat_by = {(c["tipo"], c["nome"]): c for c in cat_rows}
    sub_rows = []
    for tipo, cat_name, dre, fluxo, comp, subs in CATEGORY_DEFS:
        cat = cat_by[(tipo, cat_name)]
        for sub in subs:
            sub_rows.append({"tenant_id": TENANT_ID, "categoria_id": cat["id"], "nome": sub, "dre_grupo": dre, "natureza_fluxo_padrao": fluxo, "comportamento_padrao": comp, "ativo": True})
    subs_inserted = upsert("fin_subcategorias", sub_rows, "tenant_id,categoria_id,nome")
    centers_db = {r["nome"]: r for r in get_all("fin_centros_resultado", select="id,nome")}
    cats_db = {(r["tipo"], r["nome"]): r for r in get_all("fin_categorias", select="id,tipo,nome")}
    subs_db = {(r["categoria_id"], r["nome"]): r for r in get_all("fin_subcategorias", select="id,categoria_id,nome")}

    inserts = []
    for r, c in zip(rows, classifications):
        cat = cats_db[(c["tipo"], c["category"])]
        sub = subs_db[(cat["id"], c["subcategory"])]
        d = r["Data"]
        forma = "cartao_credito" if norm(r["Pagamento"]) == "cartao de credito" else "conta_bancaria"
        inserts.append({
            "tenant_id": TENANT_ID, "data_pagamento": d, "data_vencimento": d, "data_realizacao": d, "mes_competencia": month_start(d),
            "tipo": c["tipo"], "status": "realizado", "natureza_fluxo": c["natureza_fluxo"], "comportamento": c["comportamento"],
            "centro_resultado_id": centers_db[c["center"]]["id"], "categoria_id": cat["id"], "subcategoria_id": sub["id"], "curso_id": None,
            "forma_pagamento": forma, "banco_id": None if forma == "cartao_credito" else bank["id"], "cartao_id": card["id"] if forma == "cartao_credito" else None,
            "qtd_parcelas": 1, "descricao": clean(r["Descricao"]), "valor": str(r["Valor"]), "observacao": None,
            "origem": "importacao", "fonte_original": "Fluxo de caixa V2.xlsx", "referencia_externa": f"{SOURCE_SHEET}!{r['source_row']}", "responsavel": None,
            "classificacao_status": c["classificacao_status"],
            "metadata": {"source_file": SOURCE_FILE.name, "source_sheet": SOURCE_SHEET, "source_row": r["source_row"], "original": {k: clean(v) for k, v in r.items() if k != "Valor"}, "taxonomy_notes": c["notes"]},
        })
    inserted = []
    for i in range(0, len(inserts), BATCH_SIZE):
        inserted.extend(request("POST", "fin_lancamentos", inserts[i:i+BATCH_SIZE]))
    after = get_all("fin_lancamentos", select="id,tipo,valor,fonte_original,classificacao_status,data_pagamento")
    after_totals = defaultdict(Decimal)
    for item in after:
        after_totals[item["tipo"]] += Decimal(str(item["valor"]))
    print(json.dumps({"mode":"apply", **report, "previous_rows_snapshot": len(existing), "inserted": len(inserted), "after_rows": len(after), "after_totals": {k: str(v) for k, v in after_totals.items()}}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()







