import argparse
import json
import os
import re
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter
from datetime import datetime, date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

HML_REF = "oerdsmgiebquecqwcbox"
DEFAULT_FILE = Path(r"C:\Users\Jass_\.codex\codex-remote-attachments\01a068bf-139a-7663-af85-2472af7b8da7\B827FB91-44A1-42BA-B718-FD2BCE8E0A0C\1-Catalogo_Norwyn_Cadastro_de_Links.xlsx")
SHEET_NAME = "Cadastro de Links"
BATCH_SIZE = 100

COLUMNS = [
    "offer_name", "product_name", "offer_type", "included_products", "checkout_url", "platform",
    "price_current", "max_installments", "smart_installments", "access_time", "warranty",
    "has_coparticipation", "partner", "coparticipation_percent", "use_type", "commercial_status",
    "is_main_link", "replaced_link", "responsible", "activated_at", "notes", "campaign_name",
    "audience", "lead_origin", "special_rule",
]


def read_env():
    env = {}
    for file_name in [".env.local", ".env"]:
        p = Path(file_name)
        if not p.exists():
            continue
        for line in p.read_text(encoding="utf-8", errors="ignore").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            env.setdefault(k.strip(), v.strip().strip('"').strip("'"))
    env.update(os.environ)
    return env


def require_env():
    env = read_env()
    url = env.get("NEXT_PUBLIC_SUPABASE_URL") or env.get("SUPABASE_URL")
    key = env.get("SUPABASE_SERVICE_ROLE_KEY") or env.get("SUPABASE_SERVICE_KEY")
    if not url or not key:
        raise SystemExit("Missing Supabase URL/service key in environment.")
    if HML_REF not in url:
        raise SystemExit(f"Refusing to run: Supabase URL is not HML ref {HML_REF}.")
    return url.rstrip("/"), key


def request(method, path, body=None, params=None, prefer="return=representation,resolution=merge-duplicates"):
    base_url, key = require_env()
    endpoint = f"{base_url}/rest/v1/{path}"
    if params:
        endpoint += "?" + urllib.parse.urlencode(params, doseq=True)
    data = None if body is None else json.dumps(body, ensure_ascii=False, default=str).encode("utf-8")
    headers = {
        "apikey": key,
        "Authorization": f"Bearer {key}",
        "Content-Type": "application/json",
        "Prefer": prefer,
    }
    req = urllib.request.Request(endpoint, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=80) as resp:
            text = resp.read().decode("utf-8")
            return json.loads(text) if text else []
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="ignore")
        raise RuntimeError(f"{method} {path} failed {exc.code}: {detail}") from exc


def clean(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = " ".join(str(value).replace("\xa0", " ").strip().split())
    return text or None


def normalize(value):
    text = clean(value) or ""
    text = unicodedata.normalize("NFD", text).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def normalize_url(value):
    return (clean(value) or "").replace(" ", "")


def yes(value):
    return (clean(value) or "").lower() in {"sim", "s", "yes", "true", "1", "principal", "oficial"}


def parse_price(value):
    text = clean(value)
    if not text or text.lower() == "a confirmar":
        return None
    text = text.replace("R$", "").replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return str(Decimal(text).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
    except Exception:
        return None


def parse_int(value):
    text = clean(value)
    if not text or text.lower() == "a confirmar":
        return None
    match = re.search(r"\d+", text)
    return int(match.group(0)) if match else None


def parse_percent(value):
    text = clean(value)
    if not text:
        return None
    text = text.replace("%", "").replace(",", ".")
    try:
        return str(Decimal(text).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
    except Exception:
        return None


def offer_type(value):
    text = normalize(value)
    if "combo" in text:
        return "combo"
    if "upsell" in text:
        return "upsell"
    if "downsell" in text:
        return "downsell"
    if "bump" in text:
        return "order_bump"
    if "evento" in text:
        return "evento"
    if "especial" in text:
        return "oferta_especial"
    if "individual" in text or "produto" in text:
        return "produto_individual"
    return "outro"


def commercial_status(value):
    text = normalize(value)
    if "substit" in text:
        return "substituido"
    if "desativ" in text:
        return "desativado"
    if "paus" in text:
        return "pausado"
    if "ativo" in text:
        return "ativo"
    return "rascunho"


def parse_coparticipation(has_value, partner_value, percent_value):
    flag = normalize(has_value)
    partner = clean(partner_value)
    percent = parse_percent(percent_value)
    if flag in {"sim", "s"}:
        return True, partner, percent, None if partner and percent is not None else "coparticipacao_incompleta"
    if flag in {"nao", "no", "n"}:
        return False, None, None, None
    if partner and normalize(partner) not in {"nao", "no", "0"}:
        return True, partner, percent, None if percent is not None else "coparticipacao_percentual_a_confirmar"
    return None, None, None, "coparticipacao_a_confirmar"


def hotmart_ids(url):
    url = clean(url) or ""
    match = re.search(r"pay\.hotmart\.com\/([A-Za-z0-9]+)", url)
    offer = None
    try:
        offer = urllib.parse.parse_qs(urllib.parse.urlparse(url).query).get("off", [None])[0]
    except Exception:
        offer = None
    return match.group(1) if match else None, offer


def split_products(value):
    text = clean(value)
    if not text:
        return []
    return [part.strip() for part in text.split("|") if part.strip()]


def read_rows(path):
    wb = load_workbook(path, data_only=True)
    ws = wb[SHEET_NAME]
    rows = []
    ignored = []
    for excel_row, raw in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
        values = list(raw[:len(COLUMNS)])
        if not any(values):
            continue
        row = dict(zip(COLUMNS, values))
        offer = clean(row["offer_name"])
        url = clean(row["checkout_url"])
        if not offer and not url:
            ignored.append({"row": excel_row, "reason": "linha_sem_oferta_e_link"})
            continue
        product_name = clean(row["product_name"])
        if not product_name or not url:
            ignored.append({"row": excel_row, "reason": "produto_ou_link_ausente", "offer": offer})
            continue
        has_cp, partner, percent, cp_issue = parse_coparticipation(row["has_coparticipation"], row["partner"], row["coparticipation_percent"])
        hotmart_product, hotmart_offer = hotmart_ids(url)
        price = parse_price(row["price_current"])
        metadata_issues = []
        if cp_issue:
            metadata_issues.append(cp_issue)
        if price is None:
            metadata_issues.append("preco_a_confirmar")
        if not clean(row["use_type"]) or normalize(row["use_type"]) == "a confirmar":
            metadata_issues.append("uso_a_confirmar")
        if not clean(row["access_time"]) or normalize(row["access_time"]) == "a confirmar":
            metadata_issues.append("acesso_a_confirmar")
        rows.append({
            "excel_row": excel_row,
            "product_name": product_name,
            "product_normalized": normalize(product_name),
            "offer": {
                "name": offer or url,
                "normalized_name": normalize(offer or url),
                "offer_type": offer_type(row["offer_type"]),
                "included_products": split_products(row["included_products"]),
                "platform": clean(row["platform"]) or "Hotmart",
                "current_price": price,
                "max_installments": parse_int(row["max_installments"]),
                "smart_installments": clean(row["smart_installments"]) or "A confirmar",
                "access_time": clean(row["access_time"]) or "A confirmar",
                "warranty": clean(row["warranty"]) or "A confirmar",
                "has_coparticipation": has_cp,
                "partner": partner,
                "coparticipation_percent": percent,
                "use_type": clean(row["use_type"]) or "A confirmar",
                "commercial_status": commercial_status(row["commercial_status"]),
                "responsible": clean(row["responsible"]),
                "activated_at": None,
                "notes": clean(row["notes"]),
                "campaign_name": clean(row["campaign_name"]),
                "audience": clean(row["audience"]),
                "lead_origin": clean(row["lead_origin"]),
                "special_rule": clean(row["special_rule"]),
                "data_quality_status": "review" if metadata_issues else "trusted",
                "metadata": {"source": "Catalogo_Norwyn_Cadastro_de_Links.xlsx", "excel_row": excel_row, "normalization_issues": metadata_issues},
            },
            "link": {
                "checkout_url": url,
                "normalized_url": normalize_url(url),
                "platform": clean(row["platform"]) or "Hotmart",
                "hotmart_product_id": hotmart_product,
                "hotmart_offer_id": hotmart_offer,
                "technical_health": "nao_verificado",
                "is_main_link": yes(row["is_main_link"]),
                "attribution_confidence": "nao_atribuivel",
                "metadata": {"source": "Catalogo_Norwyn_Cadastro_de_Links.xlsx", "excel_row": excel_row, "replaced_link_text": clean(row["replaced_link"])},
            },
        })
    return rows, ignored


def get_tenant_id():
    tenants = request("GET", "tenants", params={"select": "id,nome", "nome": "ilike.*Juliana*Coutinho*", "limit": "1"})
    if not tenants:
        tenants = request("GET", "tenants", params={"select": "id,nome", "limit": "1"})
    if not tenants:
        raise SystemExit("No tenant found.")
    return tenants[0]["id"], tenants[0]["nome"]


def upsert(table, rows, conflict):
    if not rows:
        return []
    output = []
    for start in range(0, len(rows), BATCH_SIZE):
        output.extend(request("POST", table, rows[start:start+BATCH_SIZE], params={"on_conflict": conflict}))
    return output


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", default=str(DEFAULT_FILE))
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    source_path = Path(args.file)
    rows, ignored = read_rows(source_path)
    report = {
        "source_file": str(source_path),
        "valid_rows": len(rows),
        "ignored": ignored,
        "products": dict(Counter(row["product_name"] for row in rows)),
        "offer_types": dict(Counter(row["offer"]["offer_type"] for row in rows)),
        "statuses": dict(Counter(row["offer"]["commercial_status"] for row in rows)),
        "main_links": sum(1 for row in rows if row["link"]["is_main_link"]),
        "hotmart_product_ids": sum(1 for row in rows if row["link"]["hotmart_product_id"]),
        "hotmart_offer_ids": sum(1 for row in rows if row["link"]["hotmart_offer_id"]),
        "review_rows": sum(1 for row in rows if row["offer"]["data_quality_status"] == "review"),
    }
    if not args.apply:
        print(json.dumps({"mode": "dry-run", **report}, ensure_ascii=False, indent=2))
        return

    tenant_id, tenant_name = get_tenant_id()
    product_payload = [
        {"tenant_id": tenant_id, "name": product, "normalized_name": normalize(product), "status": "ativo", "metadata": {"source": source_path.name}}
        for product in sorted({row["product_name"] for row in rows})
    ]
    products = upsert("catalog_products", product_payload, "tenant_id,normalized_name")
    product_by_normalized = {item["normalized_name"]: item for item in products}

    offer_payload = []
    row_by_offer_key = {}
    for row in rows:
        product = product_by_normalized[row["product_normalized"]]
        payload = {"tenant_id": tenant_id, "product_id": product["id"], **row["offer"]}
        offer_payload.append(payload)
        row_by_offer_key[(product["id"], payload["normalized_name"])] = row
    offers = upsert("catalog_offers", offer_payload, "tenant_id,product_id,normalized_name")
    offer_by_key = {(item["product_id"], item["normalized_name"]): item for item in offers}

    link_payload = []
    for row in rows:
        product = product_by_normalized[row["product_normalized"]]
        offer = offer_by_key[(product["id"], row["offer"]["normalized_name"])]
        link_payload.append({"tenant_id": tenant_id, "product_id": product["id"], "offer_id": offer["id"], **row["link"]})
    links = upsert("catalog_sales_links", link_payload, "tenant_id,normalized_url")

    history_payload = []
    for link in links:
        offer = next(item for item in offers if item["id"] == link["offer_id"])
        history_payload.append({
            "tenant_id": tenant_id,
            "product_id": link["product_id"],
            "offer_id": link["offer_id"],
            "sales_link_id": link["id"],
            "event_type": "initial_load",
            "new_value": {"offer": offer, "link": link},
            "reason": "Carga inicial da planilha Catalogo_Norwyn_Cadastro_de_Links.xlsx",
            "actor_label": "codex_hml_initial_load",
        })
    request("POST", "catalog_link_history", history_payload, prefer="return=minimal")

    print(json.dumps({"mode": "apply", "tenant_id": tenant_id, "tenant_name": tenant_name, **report, "products_inserted_or_updated": len(products), "offers_inserted_or_updated": len(offers), "links_inserted_or_updated": len(links)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
